import re

import unicodedata

import uuid

from docx import Document
from docx.table import Table
from docx.document import Document as DocxDocument
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT

from pathlib import Path

from app.schemas.requirement import GeneratedDocument, DocumentIdentification


class DocxRenderer:
    # Limpa o nome do documento para criar o filename

    def clear_title(self, name: str) -> str:
        normalized_name = unicodedata.normalize("NFKD", name)
        without_accents = normalized_name.encode("ascii", "ignore").decode("ascii")
        lowercase = without_accents.lower()
        safe_name = re.sub(r"[^a-z0-9]+", "_", lowercase).strip("_")
        return safe_name

    # Muda o background de uma célula
    @staticmethod
    def set_cell_background(cell, color):
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        cell._tc.get_or_add_tcPr().append(shading)

    def add_identification_table(
        self,
        docx: DocxDocument,
        title: str,
        identification: DocumentIdentification | None = None,
    ) -> Table:

        fields = {
            "Versão": "version",
            "Responsável": "responsible",
            "Data": "date",
            "Status": "status",
        }

        table = docx.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.autofit = False

        header = table.rows[0]

        header.cells[0].text = "Campo"
        header.cells[1].text = "Preenchimento"

        for cell in header.cells:
            self.set_cell_background(cell, "1f4e78")

            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for paragraph in cell.paragraphs:

                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True

        row = table.add_row()

        row.cells[0].text = "Projeto / Demanda"
        row.cells[1].text = title

        for label, field in fields.items():
            row = table.add_row()
            row.cells[0].text = label

            if identification is not None:
                value = getattr(identification, field)
            else:
                value = None

            row.cells[1].text = str(value) if value is not None else "-"

        for row in table.rows[1:]:
            cell = row.cells[0]
            self.set_cell_background(cell, "D9EAF7")

            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        widths = [Cm(6), Cm(10)]

        for index, width in enumerate(widths):
            table.columns[index].width = width

        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = width

        return table

    def add_objective(self, docx: DocxDocument, objective: str | None) -> None:

        if objective and objective.strip():
            docx.add_paragraph(objective)
        else:
            docx.add_paragraph("Objetivo não informado.")

    # Adiciona uma tabela
    def add_requirements_table(
        self, docx: DocxDocument, requirements, headers, fields, widths
    ) -> Table:
        if len(headers) != len(fields) or len(headers) != len(widths):
            raise ValueError(
                "Headers, fields e widths devem possuir a mesma quantidade de elementos."
            )
        table = docx.add_table(rows=1, cols=len(headers))
        table.style = "Table Grid"
        table.autofit = False

        header_row = table.rows[0]

        for index, header_text in enumerate(headers):
            header_row.cells[index].text = header_text

        # Muda o background da célula
        for cell in header_row.cells:
            self.set_cell_background(cell, "1f4e78")

            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            for paragraph in cell.paragraphs:

                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

                for run in paragraph.runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True

        for requirement in requirements:
            row = table.add_row()

            for index, field in enumerate(fields):
                cell = row.cells[index]
                cell.text = str(getattr(requirement, field))

                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

                paragraph = cell.paragraphs[0]

                if field == "id":
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self.set_cell_background(cell, "d9eaf7")

                else:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

        for index, width in enumerate(widths):
            table.columns[index].width = width

        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = width

        return table

    # Configuração do docx
    def configure_document(self, docx: DocxDocument) -> None:
        section = docx.sections[0]
        normal_style = docx.styles["Normal"]
        title_style = docx.styles["Title"]
        heading_style = docx.styles["Heading 1"]

        normal_style.font.name = "Aptos"
        normal_style.font.size = Pt(11)

        title_style.font.name = "Aptos"
        title_style.font.size = Pt(18)
        title_style.font.bold = True
        title_style.font.color.rgb = RGBColor(31, 78, 121)

        heading_style.font.name = "Aptos"
        heading_style.font.size = Pt(14)
        heading_style.font.bold = True
        heading_style.font.color.rgb = RGBColor(31, 78, 121)

        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)

    def generate(
        self, title: str, document_id: uuid.UUID, generated_document: GeneratedDocument
    ) -> Path:

        docx = Document()

        self.configure_document(docx=docx)

        functional_requirements = generated_document.functional_requirements
        non_functional_requirements = generated_document.non_functional_requirements
        technical_restrictions = generated_document.technical_restrictions

        docx.add_heading(title, level=0)
        docx.add_heading("1. Identificação", level=1)

        self.add_identification_table(
            docx=docx, title=title, identification=generated_document.identification
        )

        docx.add_heading("2. Objetivo", level=1)

        self.add_objective(docx, objective=generated_document.objective)

        docx.add_heading("3. Requisitos Funcionais", level=1)

        # Método que gera a tabela
        if functional_requirements:
            self.add_requirements_table(
                docx=docx,
                requirements=functional_requirements,
                headers=["ID", "Requisito"],
                fields=["id", "description"],
                widths=[Cm(2.5), Cm(13.5)],
            )
        else:
            docx.add_paragraph("Nenhum requisito foi informado.")

        docx.add_heading("4. Requisitos Não Funcionais", level=1)

        if non_functional_requirements:
            self.add_requirements_table(
                docx=docx,
                requirements=non_functional_requirements,
                headers=["ID", "Requisitos Não Funcionais"],
                fields=["id", "description"],
                widths=[Cm(2.5), Cm(13.5)],
            )
        else:
            docx.add_paragraph("Nenhum requisito não funcional informado.")

        docx.add_heading("5. Restrições Técnicas", level=1)

        if technical_restrictions:

            self.add_requirements_table(
                docx=docx,
                requirements=technical_restrictions,
                headers=["ID", "Restrições Técnicas"],
                fields=["id", "description"],
                widths=[Cm(2.5), Cm(13.5)],
            )
        else:
            docx.add_paragraph("Nenhuma restrição técnica informada.")

        output_dir = Path("generated_documents")
        output_dir.mkdir(exist_ok=True)

        filename = self.clear_title(title)
        file_path = output_dir / f"{filename}_{document_id}.docx"

        docx.save(file_path)

        return file_path
