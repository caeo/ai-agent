from docx import Document
from app.schemas.requirement import GeneratedDocument


class DocxRenderer:

    def generate(self, title: str, generated_document: GeneratedDocument):

        docx = Document()

        functional_requirements = generated_document.functional_requirements
        non_functional_requirements = generated_document.non_functional_requirements
        technical_restrictions = generated_document.technical_restrictions

        docx.add_heading(title, level=0)
        docx.add_heading("1. Requisitos Funcionais", level=1)

        for requirement in functional_requirements:
            docx.add_paragraph(f"{requirement.id} - {requirement.description}")

        docx.add_heading("2. Requisitos Não Funcionais", level=1)

        for requirement in non_functional_requirements:
            docx.add_paragraph(f"{requirement.id} - {requirement.description}")

        docx.add_heading("3. Restrições Técnicas", level=1)

        for restriction in technical_restrictions:
            docx.add_paragraph(f"{restriction.id} - {restriction.description}")

        docx.save("Requisitos.docx")

        return